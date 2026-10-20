"""CZ-D3 internal publishing only. Reads saved CSV/JSON; no solver or audit import."""
from pathlib import Path
import csv, json, hashlib, math, re, sys
ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).resolve().parent
D1=ROOT/'results/CZ_D1'; D2=ROOT/'results/CZ_D2/COMMIT_B_RESUME_01'
PARENT='8147352b50168ce0f747aa2e42801bf7a2d13b35'
BANNER='内部审阅稿｜未批准对外发送｜CZ-D2完整性Gate保持FAIL'
def readcsv(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def js(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def write(p,s):p.write_text(s,encoding='utf-8')
def savecsv(name,rows):
    with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
def table(head,rows):return '| '+' | '.join(head)+' |\n|'+ '|'.join(['---']*len(head))+'|\n'+'\n'.join('| '+' | '.join(map(str,r))+' |' for r in rows)+'\n'
def fmt(v):
    x=float(v)
    return 'INF' if not math.isfinite(x) else f'{x:.3f}'
METRICS=[('range_rel','距离相对误差','%','r',100),('bearing_circular_deg','方位圆周误差','°','theta',1),('speed_rel','速度相对误差','%','v',100),('heading_circular_deg','航向圆周误差','°','psi',1)]
agg=readcsv(D2/'AGGREGATES.csv'); ai={(r['config'],r['metric']):r for r in agg}
configs=list(dict.fromkeys(r['config'] for r in agg))
diag=readcsv(D2/'OUTPUT_AND_BOUNDARY_DIAGNOSTICS.csv')
di={(r['config'],r['case_id'],r['side']):r for r in diag}
pairs=readcsv(D2/'PAIRED_COMPARISONS.csv'); stats=readcsv(D2/'STATISTICS.csv')
history=readcsv(D1/'CZ_D1_RECOMPUTED_CONFIG_SUMMARY.csv')
selected=[('H2','B_0'),('H2','U_0.05'),('AUX-D1','U0'),('AUX-D1','U5'),('T3-W1','Q1_B0'),('T3-W1','Q1_B3'),('T3-W1','Q2_B3')]
hist=[next(r for r in history if (r['experiment'],r['configuration'])==k) for k in selected]
source_paths=[D1/'CZ_D1_EVIDENCE_MASTER.csv',D1/'CZ_D1_RECOMPUTED_CONFIG_SUMMARY.csv',D1/'CZ_D1_ARCHITECTURE_METRICS.md',D1/'CZ_D1_MANEUVERS_AND_AUGMENTATIONS.md',D1/'CZ_D1_NEGATIVE_GATES_AND_GAPS.md',D1/'CZ_D1_CAP06_RECALCULATED.csv',D1/'CZ_D1_INPUT_CONTRACTS.json',D2/'AGGREGATES.csv',D2/'STATISTICS.csv',D2/'OUTPUT_AND_BOUNDARY_DIAGNOSTICS.csv',D2/'PAIRED_COMPARISONS.csv',D2/'RESOURCE_GEOMETRY.csv',D2/'COLD_AUDIT.json',D2/'AUDIT_FAILURE_SCOPE_NOTE.json',D2/'CZ_D2_FINAL_DECISION.json',D2/'PRODUCTION_PREFLIGHT.json',D2/'EXECUTION_BUDGET.json',ROOT/'results/CZ_D2/EXPERIMENT_CONTRACT.md',ROOT/'results/CZ_D2/GATE_AND_STATISTICS.md']
sources=[{'id':f'S{i+1:02d}','path':p.relative_to(ROOT).as_posix(),'sha256':sha(p),'bytes':p.stat().st_size,'snapshot_commit':PARENT,'use':'READ_ONLY_SAVED_EVIDENCE'} for i,p in enumerate(source_paths)]
sid={s['path']:s['id'] for s in sources}
savecsv('SOURCE_INDEX.csv',sources)
write(OUT/'INTERNAL_AUTHORIZATION.json',json.dumps({'stage':'CZ-D3_INTERNAL_REPORT','parent':PARENT,'branch':'research/cz-d3','authority':'Latest research lead review in conversation; source is user-supplied audit, not a new Codex full audit','external_release_authorized':False,'new_scientific_execution_authorized':False,'original_integrity_pass':False,'accepted_interpretation':'NO_SUCCESSFUL_PRIVATE_READ_OBSERVED_IN_RECORDED_TRACE','not_certified':'PRIVATE_DATA_ISOLATION_FULLY_CERTIFIED','historical_gate_changes':False,'R4_percent':0},indent=2)+'\n')
conditions={}
for c in configs:
    q,b,p=c.split('_');q1=q=='Q1'
    conditions[c]={'nodes':2,'baseline_km':5 if b=='B0' else math.sqrt(234),'offset_km':'(0,+/-5)' if b=='B0' else '(3,+/-15)','bearing_sigma_deg':.05 if q1 else .075,'navigation_sigma_m_each_axis_epoch':25 if q1 else 37.5,'generating_common_bias_uniform_bound_deg':.05 if q1 else .075,'generating_differential_half_bias_uniform_bound_deg':.025 if q1 else .0375,'estimator_bias_support_multiplier':1 if p=='P0' else 2,'time_s':1200,'epochs':121,'turn_at_s':600,'turn_deg':15,'panel_states':12,'draws_per_state':500,'draws_per_side_per_state':250,'config_draws':6000,'independent_base_draws_campaign':6000,'paired_across_all_configurations':True,'depth':'NOT_EVALUATED','signed_bearing_sync_association':'ASSUMED_NOT_EXTRACTED','measurement_chain':'SYNTHETIC_FEATURE_MODEL','integrity_gate':'FAIL_RETAINED'}
write(OUT/'OBSERVATION_CONDITIONS.json',json.dumps(conditions,indent=2)+'\n')
precision=[]
for r in hist:
    for _,name,unit,h,factor in METRICS:
        precision.append({'study':r['experiment'],'configuration':r['configuration'],'parameter':name,'unit':unit,'P50_mean_scene':r[h+'_unconditional_mean_scene_P50'],'P50_pooled':'NOT_REPRODUCED_HERE','P95_worst_scene':r[h+'_unconditional_worst_scene_P95'],'P99_worst_scene':'SEE_SOURCE_SCENE_TABLE','accepted':r['accepted'],'planned':r['n'],'nonoutput_rate':1-int(r['accepted'])/int(r['n']),'geometries':r['scenes'],'depth':'NOT_EVALUATED','evidence_scope':'CZ_D1_SCOPED_HISTORICAL_CONDITIONAL','source':str((D1/'CZ_D1_RECOMPUTED_CONFIG_SUMMARY.csv').relative_to(ROOT)).replace('\\','/'),'selector':r['experiment']+'/'+r['configuration'],'conditions':'../CZ_D1/CZ_D1_INPUT_CONTRACTS.json'})
for c in configs:
    d=di[c,'POOLED','BOTH']
    for m,name,unit,h,factor in METRICS:
        a=ai[c,m]
        precision.append({'study':'CZ-D2','configuration':c,'parameter':name,'unit':unit,'P50_mean_scene':float(a['mean_of_12_scene_P50'])*factor,'P50_pooled':float(a['pooled_P50'])*factor,'P95_worst_scene':float(a['worst_scene_P95'])*factor,'P99_worst_scene':float(a['worst_scene_P99'])*factor,'accepted':d['accepted'],'planned':d['planned'],'nonoutput_rate':d['failed_rate'],'geometries':12,'depth':'NOT_EVALUATED','evidence_scope':'CONDITIONAL_SIMULATION_WITH_INTEGRITY_EXCEPTION','source':'results/CZ_D2/COMMIT_B_RESUME_01/AGGREGATES.csv','selector':c+'/'+m,'conditions':'OBSERVATION_CONDITIONS.json#'+c})
savecsv('PRECISION_OBSERVATION_CONDITIONS.csv',precision)
savecsv('D2_BOUNDARY_AND_TAIL.csv',[r for r in diag if r['case_id']=='POOLED'])
savecsv('D2_PAIRED_CHANGES.csv',[r for r in pairs if r['case_id']=='POOLED'])
# Copy all relevant final reported quantities for side/scenario transparency; no estimator rerun.
savecsv('D2_SCENE_AND_SIDE_QUANTILES.csv',[r for r in stats if r['metric'] in [m[0] for m in METRICS]])
resource=readcsv(D2/'RESOURCE_GEOMETRY.csv');res=[]
for b in ['B0','B3']:
    rr=[r for r in resource if r['quality']=='Q1' and r['geometry']==b]
    res.append({'geometry':b,'unique_case_side_trajectories':len(rr),'baseline_km':rr[0]['node_baseline_km'],'main_min_km':min(float(r['main_min_km']) for r in rr),'main_max_km':max(float(r['main_max_km']) for r in rr),'aux_min_km':min(float(r['aux_min_km']) for r in rr),'aux_max_km':max(float(r['aux_max_km']) for r in rr),'main_outside45_60_epochs':sum(int(r['main_outside45_60']) for r in rr),'aux_outside45_60_epochs':sum(int(r['aux_outside45_60']) for r in rr),'epochs_denominator':len(rr)*121,'scope':'GEOMETRY_NOT_DETECTION; Q1/Q2 identical nominal tracks counted once','source':'results/CZ_D2/COMMIT_B_RESUME_01/RESOURCE_GEOMETRY.csv'})
savecsv('DEPLOYMENT_GEOMETRY.csv',res)
arch=[
['A0','原8元、14 m HLA','原资源','有符号方位的条件基线；多频有限候选机制','连续全域支持、实际接收可靠性、五维精度未建立','原始同步复数据、阵形和通道标定、左右舷与干扰处理','S01/S02/S03'],
['A1','单HLA＋机动','0/2/5/10/15°与1200 s航迹','特定RC2+RC3 15°例距离候选跨度4 km；RC2-only仍15 km','候选跨度不是P95；多数其它方案已含15°转向','可实施机动、拖阵姿态与导航误差；需同合同配对证据','S04'],
['A2','单HLA＋附加声学特征','三频或六窗径向特征假设','H1公共源幅剖面机制；H2速度中位收益','离网格搜索、径向映射、频率零点、接收提取与尾部缺口','可跟踪谱线、源频参考、符号/零点、增益及传播映射','S03/S05'],
['A3','HLA＋短垂向','增加4通道、4 m垂向跨度','局部深度信息比2.3407–4.4743（相对等通道共深对照）','不是米级深度精度；宽水平支持与环境影响未闭合','同步、差分增益、阵形、可信传播与实际相对观测','S03'],
['A4','整体倾斜8元','原通道，姿态控制','前向精度诊断和误差归因','性能NOT_EVALUATED；不能判阵型物理无效','可信相干前向及姿态合同；旧T2/FUNC2链保持关闭','S03/S05'],
['A5','HLA＋辅助节点','第二接收节点；5–15.297 km','AUX/T3条件四参数；D2有完整性例外的泛化结果','方位尾部非统一改善；硬偏置支持；无深度与实收认证','部署、导航、同步、同目标关联、可用有符号测向、偏置依据','S02/S08/S18'],
['A6','联合增强','仅计实际已做组合','已测试的转向＋辅助列于各自合同','CAP06 P3未执行；不能拼接H3与T3为五维性能','逐观测准入后取得同批联合输入；无已验证Pareto前沿','S03/S05']]
savecsv('ARCHITECTURE_AND_RESOURCE_MATRIX.csv',[dict(zip(['id','architecture','resource','evidence','limit','required_conditions','source_ids'],r)) for r in arch])
claims=[
('C01','原单HLA五维联合精度尚未建立','S01/S03','模型、候选、经验误差和局部信息分栏，不合成系统指标'),
('C02','H2 U_0.05速度场景等权中位0.7538%；最坏几何无条件P95为INF','S02/S03','假设径向量；250/256输出；不是实际提取性能'),
('C03','H3局部深度信息比2.3407–4.4743','S03','条件局部信息，不换算米或误差倍数'),
('C04','CAP06 CAL20覆盖21378/22648=94.3924%，冻结门95%未过','S03/S06','目标存在历元角度候选覆盖，不是目标四维精度或独立航迹覆盖'),
('C05','D2 Q1/P0 B3速度最坏场景P95约11.557%，B0约30.932%','S08/S09','新12几何条件仿真，完整性FAIL；非全部连续域保证'),
('C06','D2 Q1/P0 B3方位尾部略差','S08/S09','0.143109°对0.133745°；不存在四参数统一改善'),
('C07','B3各配置6000/6000至少一个偏置触界','S10','接受表示收敛有限；不表示正确或校准恢复'),
('C08','P1仅扩大估计支持，仍有硬先验，生成偏置不变','S18/S19','不称超界失配、无先验或纯物理信息变化'),
('C09','539763 PASS、1 FAIL保持；没有整体科学PASS','S13/S14/S15','日志未观察到成功私有读取，不是OS隔离认证'),
('C10','旧Q1/B3速度8.5659%与新11.557%不同','S02/S08','不同目标面板的泛化差异，不能作配对改进/退化实验'),
('C11','NUM03名义首折返几何，不是实海域探测覆盖','S05','单一WOA来源族；相干提供器与能量归因限制仍在'),
('C12','R3最终双候选只差4°/5°航向','S03/S04','114576→2仅机制，不是连续零误差'),
('C13','新增资源对应的是条件证据，未形成工程最低配置或价格','S03/S04/S18','不提供未验证的硬件能力或成本数字')]
savecsv('CLAIM_EVIDENCE_LEDGER.csv',[{'claim_id':c,'statement':s,'source_ids':i,'mandatory_qualification':q,'release_status':'INTERNAL_PENDING_LEAD_REVIEW'} for c,s,i,q in claims])
# Tables rendered from source cells, with explicit aggregation names.
hrows=[]
for r in hist:
    hrows.append([r['experiment']+' '+r['configuration']]+[fmt(r[h+'_unconditional_mean_scene_P50'])+' / '+fmt(r[h+'_unconditional_worst_scene_P95']) for _,_,_,h,_ in METRICS]+[r['accepted']+'/'+r['n']])
drows=[];medrows=[];tailrows=[]
for c in configs:
    drows.append([c]+[fmt(float(ai[c,m]['worst_scene_P95'])*f) for m,_,_,_,f in METRICS])
    medrows.append([c]+[fmt(float(ai[c,m]['mean_of_12_scene_P50'])*f)+' / '+fmt(float(ai[c,m]['pooled_P50'])*f) for m,_,_,_,f in METRICS])
    tailrows.append([c]+[fmt(float(ai[c,m]['worst_scene_P99'])*f) for m,_,_,_,f in METRICS])
heads=['配置','距离 %','方位 °','速度 %','航向 °']
brows=[[c,di[c,'POOLED','BOTH']['accepted']+'/6000',di[c,'POOLED','BOTH']['bias_at_bound_count'],di[c,'POOLED','BOTH']['state_at_bound_count'],di[c,'POOLED','BOTH']['outside_original_bias_count']] for c in configs]
pairrows=[]
for r in pairs:
    if r['baseline']=='Q1_B0_P0' and r['candidate']=='Q1_B3_P0' and r['case_id']=='POOLED' and r['side']=='BOTH' and r['metric'] in [m[0] for m in METRICS]:
        pairrows.append([next(m[1] for m in METRICS if m[0]==r['metric']),r['better'],r['worse'],r['equal'],r['paired_executed']])
checks=[]
for c in configs:
    for m,_,_,_,_ in METRICS:
        vv=[float(r['P95']) for r in stats if r['config']==c and r['metric']==m and r['mode']=='unconditional' and r['side']=='BOTH' and r['case_id']!='POOLED']
        checks.append({'check':'saved_table_max_identity','config':c,'metric':m,'pass':len(vv)==12 and max(vv)==float(ai[c,m]['worst_scene_P95'])})
assert all(x['pass'] for x in checks)
write(OUT/'TABLE_CHECKS.json',json.dumps({'scope':'DOCUMENT_TABLE_TRANSCRIPTION_ONLY_NOT_COLD_AUDIT','checks':checks,'scientific_gate_changed':False},indent=2)+'\n')
# Scientific plots use ONLY archived aggregate/statistic tables.
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
font_manager.fontManager.addfont('C:/Windows/Fonts/msyh.ttc')
plt.rcParams.update({'font.family':'Microsoft YaHei','axes.spines.top':False,'axes.spines.right':False,'font.size':10,'pdf.fonttype':42,'svg.fonttype':'none'})
FIG=OUT/'figures';FIG.mkdir(exist_ok=True)
navy='#15364C';teal='#158078';orange='#C46B32';red='#B64545'
def savefig(fig,name,caption):
    fig.text(.5,.015,caption,ha='center',fontsize=9,color='#555555')
    for ext in ['png','pdf','svg']:fig.savefig(FIG/(name+'.'+ext),dpi=180,bbox_inches='tight')
    plt.close(fig)
fig,axs=plt.subplots(2,2,figsize=(10.8,6.4));fig.subplots_adjust(hspace=.48,bottom=.12,top=.89)
for ax,(m,n,u,_,f) in zip(axs.ravel(),METRICS):
    vals=[float(ai[c,m]['worst_scene_P95'])*f for c in ['Q1_B0_P0','Q1_B3_P0']]
    ax.bar(['B0 · 5 km','B3 · 15.297 km'],vals,color=[navy,teal],width=.5)
    ax.set_title(n+'（'+u+'）');ax.set_ylim(0,max(vals)*1.32)
    for i,v in enumerate(vals):ax.text(i,v*1.04,fmt(v),ha='center')
fig.suptitle('CZ-D2｜Q1/P0：宽基线改善三个参数，方位尾部略差',fontsize=15)
savefig(fig,'F01_D2_FOUR_PARAMETER_P95','内部稿 · 完整性Gate FAIL · 12几何最坏无条件经验P95 · 来源S08/S09 · 非工程指标')
fig,axs=plt.subplots(1,3,figsize=(11,4.5));fig.subplots_adjust(bottom=.22,top=.8,wspace=.35)
for ax,(m,n,u,_,f) in zip(axs,[METRICS[0],METRICS[1],METRICS[2]]):
    for q,col in [('Q1',teal),('Q2',orange)]:
        vv=[float(ai[f'{q}_B3_{p}',m]['worst_scene_P95'])*f for p in ['P0','P1']]
        ax.plot([0,1],vv,'o-',color=col,label=q)
        for i,v in enumerate(vv):ax.annotate(fmt(v),(i,v),xytext=(0,7),textcoords='offset points',ha='center',fontsize=9)
    ax.set_xticks([0,1],['P0原支持','P1两倍支持']);ax.set_title(n+'（'+u+'）');ax.margins(x=.3,y=.35);ax.legend(frameon=False)
fig.suptitle('B3偏置支持敏感性：生成偏置不变，估计边界扩大',fontsize=15)
savefig(fig,'F02_BIAS_SUPPORT_SENSITIVITY','内部稿 · Gate FAIL · 各组B3均6000/6000偏置触界 · P1仍有硬先验 · 来源S08/S10/S18')
fig,ax=plt.subplots(figsize=(10.5,4.8));xs=list(range(8));ww=.25
for j,(key,label,col) in enumerate([('pooled_P50','池化P50',teal),('worst_scene_P95','最坏几何P95',navy),('worst_scene_P99','最坏几何P99',orange)]):
    vals=[float(ai[c,'speed_rel'][key])*100 for c in configs]
    ax.bar([x+(j-1)*ww for x in xs],vals,ww,label=label,color=col)
ax.set_xticks(xs,[c.replace('_','\n') for c in configs]);ax.set_ylabel('速度相对误差（%）');ax.legend(frameon=False,ncol=3);ax.set_title('速度中位与尾部：统计层级分别展示',fontsize=15);fig.subplots_adjust(bottom=.24)
savefig(fig,'F03_SPEED_MEDIAN_AND_TAIL','内部稿 · Gate FAIL · P50与最坏几何尾部不是同一分布的三个分位数 · 来源S08')
fig,ax=plt.subplots(figsize=(9.7,4.3));yy=list(range(len(pairrows)))
a=[int(r[1]) for r in pairrows];b=[int(r[2]) for r in pairrows];e=[int(r[3]) for r in pairrows]
ax.barh(yy,a,color=teal,label='B3误差更小');ax.barh(yy,b,left=a,color=orange,label='B3误差更大');ax.barh(yy,e,left=[a[i]+b[i] for i in yy],color='#C7CFD2',label='相等')
ax.set_yticks(yy,[r[0] for r in pairrows]);ax.set_xlabel('配对记录数（每参数6000；四参数不是独立重复）');ax.set_xlim(0,6000)
for i in yy:
    if a[i]:ax.text(a[i]/2,i,str(a[i]),ha='center',va='center',color='white')
    if b[i]:ax.text(a[i]+b[i]/2,i,str(b[i]),ha='center',va='center',color='white')
ax.legend(ncol=3,loc='upper center',bbox_to_anchor=(.5,1.18),frameon=False);fig.subplots_adjust(bottom=.23,top=.78)
savefig(fig,'F04_PAIRED_CHANGES','内部稿 · Q1/P0 · 完整性Gate FAIL · 来源S11 · 配对计数不证明总体改善概率')
# Internal report sections. No performance inference beyond saved records.
sections=[]
def sec(title,text):sections.append((title,text.strip()))
sec('1. 交付结论与审阅范围','''本稿用于716项目首次阶段性交付的内部审阅，尚未批准向甲方发送或作为性能验收文件。研究目标是“输入观测条件与精度 → 五维状态输出精度 → 适用边界与资源需求”。当前不能承诺原单平台单HLA的五维联合精度。

原单HLA是主线。已有成果包括机动与多频候选区分、源未知量消除机制、假设径向特征对速度的潜力，以及接收提取的具体失败边界。增强方案作为条件参照：短垂向观测有局部深度信息增量，双节点有四参数经验误差分布；它们尚不能组成已验证的五维系统。

CZ-D1限定范围归并结论已获负责人接受。CZ-D2完成48,000项计算，保留较强数值一致性证据，但冻结完整性Gate仍为FAIL，整体科研PASS不予授予。负责人接受标准库缓存访问解释并允许本轮内部编稿，不构成原Gate追认。所有CZ-D2精度均按“有完整数值证据、存在完整性例外的条件性仿真结果”引用。

本轮仅读取保存的表格与文档、进行展示转换和表间一致性检查。不重新访问并复算全部D盘原始观测，不运行传播、优化、Monte Carlo、Fisher或冷审计。main、历史Gate和R4=0%保持。''')
sec('2. 指标定义与证据等级','''状态顺序为距离r、方位θ、深度z、速度v、航向ψ。相对误差逐样本除以真实r或v；角误差为最小圆周差（度）。深度未评估写NOT_EVALUATED，不填0。拒判和数值失败计INF；未执行与失败分开。分位数继承nearest-rank，不对INF插值。

证据分为：①解析或局部信息；②有限离散候选；③假设观测下的连续估计误差；④合成接收数据实际提取；⑤真实设备/海域验证。类别不同，不构成可直接排序的精度排行榜。已有材料未形成第⑤类工程能力认证。

统一对照表主P50为各几何P50的等权算术均值，P95为最坏几何无条件P95；另列CZ-D2池化P50。二者不是同一分布的分位数对。最坏场景可随参数和配置变化；两个最坏P95相减不是逐样本配对误差差的P95。

H2每几何32次、AUX-D1每几何128次、T3-W1及CZ-D2每几何500次。有限面板P95均为探索性经验统计，没有总体95%置信保证。配置和历元不能当作独立航迹；CZ-D2的48,000项配置来自6,000组配对基础创新。''')
sec('3. 原单HLA及机动：可以交付的进展','''原阵名义8元、2 m间距、14 m孔径。AUX-D1 U0与H2 B分别提供各自合同内的方位基线，并不穷尽原阵声学能力，也不是原阵物理极限。

历史R3同合同转角0/2/5/10/15°对照中，MAIN、源深180 m、τ=0.5 dB的RC2-only距离跨度均为15 km；结合RC3后，仅15°实例收至4 km。该数值是有限候选跨度，不是连续测距误差。多数后续实验本已采用600 s后15°转向，不能将整套增强收益归功于机动。

最终114,576→2标签是理论机制验证，最终两候选仅航向4°/5°不同，不作为主要性能成果或连续零误差。未建立转角、时长到四参数P95的完整函数。

H1证明公共源幅剖面化的条件价值，但冻结离网格声学搜索没有建立可接受解；未找到解不能当作物理信息不存在。H2主U条件注入六窗径向幅值（σ=0.05 m/s），速度场景等权中位误差约0.7538%，最坏几何无条件P95仍为INF。原始声压到该特征的映射未准入，不能把注入精度当作设备能力。[S02–S05]''')
sec('4. 历史条件精度：保留收益与失败分母',table(['实验/配置','距离 P50/P95 %','方位 P50/P95 °','速度 P50/P95 %','航向 P50/P95 °','输出/计划'],hrows)+'''
P50为场景等权中位误差汇总，P95为最坏场景无条件P95。各行深度NOT_EVALUATED。各实验有不同目标面板、估计器与误差合同，只能在同实验内作资源对照；禁止跨行解释为单因素算法优劣。[S02/S03]

H2：方位σ=0.1°、精确导航、有符号方位、1200 s及六个200 s径向窗；U主档零点已校准但接近/远离符号未知。250/256输出不是250次精度达标。
AUX-D1：U0/U5配对同一主节点观测；随机方位0.05°、导航25 m/轴、共模±0.05°、半差±0.025°及布放偏角±10°，四状态稳健处理器不估固定偏置。
T3-W1：B0和B3都是双节点，使用偏置有界估计，不能冒充原单HLA对照。Q1/Q2定义见下一节；所有历史数字仅为其冻结模型条件。''')
sec('5. 深度、倾斜、环境与接收链边界','''H3短垂向B1相对等通道共深B2的局部深度信息比为2.3407–4.4743，依赖名义环境、校准和局部水平状态剖面。不能将信息比转为误差改善倍数或米级深度P95。宽水平多分支、环境传播及实际提取仍是必要关口。

T2倾斜阵及短垂向阵的定位性能未准入；前向数值失败与理想波导精度归因不等于阵型物理无效。旧T2/FUNC2修补链保持关闭，本稿不提出继续加密或补跑。

NUM03支持单一WOA来源族下的名义第一深折返几何，两种SSP表示对照不等于真实海域传播稳健性。非相干能量、相干模态、可检测性与工程覆盖必须分别论证。CZO04揭示每历元自由源级可吸收单频强度变化，当前单频包络不构成一般可靠距离锚。

CAP06 CAL/20 dB主条件：21,378/22,648目标存在历元被角度候选覆盖，94.3924%低于冻结95%门；这不是四维目标精度或独立航迹成功率。未校准10/20 dB无目标分量提取输出，O2前向未准入，P3融合未执行。不得用较易子集替代主条件。其原始哈希与冷检查属于历史执行证据，本稿未再做全量复算。[S03/S05/S06]''')
sec('6. CZ-D2观测合同与独立样本口径','''B0辅助偏移(0,±5) km；B3为(3,±15) km，节点间距15.297 km。两节点保持全局固定偏移共速平移，主平台2 m/s、600 s后转15°，1200 s/121历元。两侧布放分层不等于原HLA左右舷消歧。

Q1：每节点随机方位σ=0.05°，导航每轴每历元25 m；固定共同偏置c~Uniform±0.05°，差模半差d~Uniform±0.025°。Q2相应为0.075°、37.5 m、±0.075°、±0.0375°；Q2同时改变多个条件，不能归因于单一测向精度。

主/辅偏置为c−d与c+d。P0按原支持估计，P1仅将两偏置估计支持扩大2倍；生成偏置仍处P0范围。P1仍是有先验估计，不是实际超界偏置失配试验。归一化尺度变化影响TRF数值轨迹，结果称“冻结处理器的偏置支持敏感性”。

新面板12个预注册离网格状态，每状态500次、每侧250次；6,000组基础创新配对8配置，24,000条公开资源记录，48,000终点。它不是保密外部盲测集。生成代码在面板物化前冻结，V3排除762条身份按已审计范围绑定。

物理状态域r∈[45,60] km、θ∈[−5,5]°、v∈[1,3] m/s、ψ∈[−15,15]°。估计器仅读取公开时间、报告位置与有符号方位，单次局部TRF终点；没有全局最优和连续分支覆盖证书。实际同步、目标关联、有符号测向与偏置范围均为假设。[S18/S19]''')
sec('7. CZ-D2四参数经验P95：完整八组配置',table(heads,drows)+'''
以上为12几何最坏无条件经验P95，全部深度NOT_EVALUATED；原冻结完整性Gate FAIL。每配置6,000/6,000返回满足数值接受条件的终点，接受率100%不表示定位正确率100%。[S08–S10]

Q1/P0：B3距离、速度、航向尾部优于B0，方位略差。速度最坏P95由30.9322%降至11.5570%，两最坏统计量相对减少约62.6%，不是同一记录的改善百分比。旧T3-W1 Q1/B3为8.5659%，新面板未达到该数字，不能据此承诺全部适用几何速度误差低于10%。

Q2/P0的B3速度最坏P95约17.426%。本次表间抄录检查使用档案原值；用户审阅给出的四舍五入末位与表格显示可能不同，精确值以CSV为准。没有把5%或10%设为甲方验收线。

![四参数P95](figures/F01_D2_FOUR_PARAMETER_P95.png)''')
sec('8. 中位、尾部、分侧和配对结果','各格为“场景等权P50 / 池化P50”，单位同列标题。[S08]\n\n'+table(heads,medrows)+'\n最坏场景无条件P99（不是总体置信上界）：\n\n'+table(heads,tailrows)+'\nQ1/P0、B3相对B0逐记录配对：\n\n'+table(['参数','误差更小','误差更大','相等','配对分母'],pairrows)+'''
全部逐几何、正负侧、池化、条件/无条件P50/P90/P95/P99见D2_SCENE_AND_SIDE_QUANTILES.csv；所有配置配对与失败转移见D2_PAIRED_CHANGES.csv。两侧各250/场景、合计各3,000/配置，不能把两侧当两批独立目标。单参数更好记录数不代表四参数联合准确率。

![速度尾部](figures/F03_SPEED_MEDIAN_AND_TAIL.png)

![配对变化](figures/F04_PAIRED_CHANGES.png)''')
sec('9. 偏置触界、物理边界与估计风险',table(['配置','数值接受/计划','至少一偏置触界','至少一物理状态触界','超过原偏置支持'],brows)+'''
触界沿用归一化tol=1e−6，P1尺度扩大使其物理容差相应变化。P1超过原P0支持不等于超出当前估计域，也不是生成器发生超界。物理状态触界、单局部终点和极高偏置触界率都限制了可交付措辞。

Q1/B3放宽支持后，距离P95约0.684%→1.018%，方位0.143°→0.218°，速度11.557%→11.706%。数据似然在同一物理状态下不变，不能把所有差异都归于物理信息减少。无硬支持的局部Fisher与有界Monte Carlo不是同一统计对象；不把局部1.96σ当经验P95或将Uniform改为Gaussian先验。

D2_BOUNDARY_AND_TAIL.csv保留逐侧和池化的大误差计数：距离10%、方位0.5°、速度20%、航向5°仅为原描述性尾部标尺，不是接受条件或工程要求。[S10/S18/S19]

![偏置支持](figures/F02_BIAS_SUPPORT_SENSITIVITY.png)''')
sec('10. 完整性例外必须与数值结果同时披露','''正式主分析一次、独立冷审计一次、重试零；48,000/48,000配置完成。冷审计539,763 PASS、1 FAIL，返回非零状态，integrity_pass=false保持。32项四参数最坏P95的表间一致性不能代替该冷审计，更不能抵消失败。

唯一失败private_read_isolation的冻结判据要求denied==[]，并要求成功读取不含私有/panel路径。日志拒绝zipfile与encodings.cp437两份标准库.pyc缓存，随后成功读取对应源代码。记录中没有成功私有真值读取路径；负责人接受缓存回退解释。

可使用的表述：NO_SUCCESSFUL_PRIVATE_READ_OBSERVED_IN_RECORDED_TRACE。不能使用：PRIVATE_DATA_ISOLATION_FULLY_CERTIFIED。Python审计钩子不是操作系统级沙箱，拒绝日志没有完整打开标志。整体科学PASS不授予。

本稿不修改原审计、代码或STOP，不重跑48,000配置，也不建议以修审计后重跑取得PASS。早期H2严格相等FAIL、REPAIR的671 PASS/1 FAIL及14项补充检查同样保持原身份。新的负责人裁决单独保存，不追溯覆盖历史。[S13–S17]''')
sec('11. 资源需求与适用范围',table(['架构','资源变化','已有证据','仍缺条件'],[[r[0]+' '+r[1],r[2],r[3],r[5]] for r in arch])+'''\n资源表描述观测必要条件与已测试名义配置，不是最低装备数量、价格、可达精度保证或已验证Pareto边界。原HLA缺失的信号信息不能仅靠算法承诺补回。\n\n'''+table(['几何','节点间距km','辅助距离最小/最大km','辅助域外历元/2904'],[[r['geometry'],fmt(r['baseline_km']),fmt(r['aux_min_km'])+' / '+fmt(r['aux_max_km']),str(r['aux_outside45_60_epochs'])+'/2904'] for r in res])+'''
几何统计取24条名义目标—侧向组合轨迹，每条121历元，Q1/Q2相同名义轨迹只计一次。45–60 km只是历史研究范围，不能证明第一CZ探测覆盖；域外值保留，不删不换面板。实际15 km级部署、测向能力和海域覆盖仍需独立证据。[S12/S18]

后续资料优先级：①原8元同步复基带、时间戳、AGC和阵元标定；②拖阵姿态、阵形与导航协方差；③谱线连续性、源频参考和干扰分量；④SSP/海深/底质及变化范围；⑤辅助节点同步、同目标关联、左右舷判别与偏置标定；⑥短垂向附加通道的差分增益与部署条件。当前均不能从模拟参数反推设备已经具备能力。''')
sec('12. 对外措辞清单与收口决定','''可写：原单HLA完成多条条件机制及能力边界研究；给出假设观测到状态误差的可追溯证据，深度和实际接收链仍需验证。可写：增加辅助节点在给定测向、导航和偏置支持下改善部分水平参数，CZ-D2保留完整性例外。

不可写：单HLA已实现五维联合小于10%；100%接受等于100%定位正确；局部深度信息增益等于米级测深；B3对全部参数均更优；P1验证任意超界偏置鲁棒；9个名义环境单元等于9个海域；通过修正说明将原冷审计改成全PASS。

申报研究问题应落在“第一CZ条件下，可提取观测如何解除状态与源、增益、偏置、环境之间的退化，以及代价是什么”。已有正、负证据均可支持问题定义，尚不能构成工程性能验收。

本轮完成内部技术报告、图件、条件表和PPT草稿后STOP。下一步仅等待负责人审查指标措辞、证据等级和例外披露；未授权发送甲方，未授权新实验，CZ-D2数值路线收口。''')
source_table=table(['编号','固定快照来源'],[[s['id'],'['+s['path'].split('/')[-1]+'](../../'+s['path'].split('results/',1)[-1]+')'] for s in sources])
# Correct links from results/CZ_D3 to sibling result directories.
source_table=source_table.replace('(../../','(../')
report='# 716第一会聚区多参数定位：首次阶段性交付内部技术报告\n\n'+BANNER+'\n\n基于CZ-D2提交 `'+PARENT+'`；历史CZ-D1最终提交 `a580499475cc4148f1195e6256c4813c6e225726`。\n\n'+'\n\n'.join('## '+t+'\n\n'+body for t,body in sections)+'\n\n## 附录：来源索引\n\n'+source_table+'\n全部路径、SHA256、固定快照与用途见SOURCE_INDEX.csv；逐项措辞见CLAIM_EVIDENCE_LEDGER.csv。\n'
write(OUT/'CZ_D3_INTERNAL_TECHNICAL_REPORT.md',report)
# DOCX with editable text/tables and embedded source-derived figures.
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
doc=Document();normal=doc.styles['Normal'];normal.font.name='Microsoft YaHei';normal.font.size=Pt(10)
normal.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
for style in ['Title','Heading 1','Heading 2']:
    doc.styles[style].font.name='Microsoft YaHei';doc.styles[style].element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
doc.sections[0].header.paragraphs[0].text=BANNER
doc.sections[0].footer.paragraphs[0].text='CZ-D3内部审阅稿｜R4=0%｜全部精度须连同条件与例外引用'
lines=report.splitlines();i=0
while i<len(lines):
    line=lines[i]
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            if not re.match(r'^\|[\- :|]+\|$',lines[i]):rows.append([x.strip() for x in lines[i].strip('|').split('|')])
            i+=1
        tb=doc.add_table(rows=0, cols=len(rows[0]));tb.style='Light Shading Accent 1'
        for row in rows:
            cells=tb.add_row().cells
            for cell,txt in zip(cells,row):cell.text=re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',txt)
        for rr in tb.rows:
            for cc in rr.cells:
                for pp in cc.paragraphs:
                    for run in pp.runs:run.font.size=Pt(8)
        continue
    if line.startswith('!['):
        fn=re.search(r'\(([^)]+)\)',line).group(1);doc.add_picture(str(OUT/fn),width=Inches(6.1))
    elif line.startswith('# '):doc.add_heading(line[2:],0)
    elif line.startswith('## '):doc.add_heading(line[3:],1)
    elif line.strip():doc.add_paragraph(line.replace('`','').replace('**',''))
    i+=1
doc.save(OUT/'CZ_D3_INTERNAL_TECHNICAL_REPORT.docx')
# Presentation: editable text/tables, no external publishing. Speaker notes cite fixed sources.
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor
prs=Presentation();prs.slide_width=PI(13.333);prs.slide_height=PI(7.5)
slide_index=[]
def color(h):return RGBColor.from_string(h.strip('#'))
def textbox(sl,x,y,w,h,text,size=22,col=navy,bold=False):
    sh=sl.shapes.add_textbox(PI(x),PI(y),PI(w),PI(h));tf=sh.text_frame;tf.word_wrap=True
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Microsoft YaHei';p.font.size=PP(size);p.font.bold=bold;p.font.color.rgb=color(col);p.space_after=PP(12)
    return sh

def slide(title,subtitle='',source='S01/S03',notes=''):
    sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=color('#F5F7F8')
    textbox(sl,.45,.25,12.4,.65,title,28,bold=True)
    if subtitle:textbox(sl,.5,1,12.1,.65,subtitle,16,col=teal)
    textbox(sl,.5,7.03,11.9,.25,BANNER+'  |  '+source,9,col='#695651')
    textbox(sl,12.3,7.03,.6,.25,str(len(prs.slides)),10)
    sl.notes_slide.notes_text_frame.text='来源：'+source+'；SOURCE_INDEX.csv提供固定SHA256。\n'+notes+'\n内部稿，未批准外发。原完整性FAIL保持。'
    slide_index.append({'slide':len(prs.slides),'title':title,'source_ids':source,'qualification':notes})
    return sl

def bullets(title,items,source='S01/S03',subtitle='',notes=''):
    sl=slide(title,subtitle,source,notes)
    for i,item in enumerate(items):textbox(sl,.7,1.85+i*.97,11.9,.9,'• '+item,21)
    return sl

def ptable(sl,headers,rows,y=1.85,colwidths=None,font=17):
    sh=sl.shapes.add_table(len(rows)+1,len(headers),PI(.6),PI(y),PI(12.1),PI(min(4.8,.48*(len(rows)+1))));tb=sh.table
    if colwidths:
        for c,w in zip(tb.columns,colwidths):c.width=PI(w)
    for i,row in enumerate([headers]+rows):
        for j,text in enumerate(row):
            cell=tb.cell(i,j);cell.text=str(text);cell.fill.solid();cell.fill.fore_color.rgb=color(navy if i==0 else ('#E9F0F2' if i%2 else '#FFFFFF'))
            for p in cell.text_frame.paragraphs:
                p.font.name='Microsoft YaHei';p.font.size=PP(font);p.font.color.rgb=color('#FFFFFF' if i==0 else navy)
    return sh
bullets('716首次阶段性交付｜内部审阅稿',['输入观测条件与精度 → 目标状态精度 → 适用边界与资源需求','原单HLA为主线；辅助节点结果作为条件性增强参照','深度、真实接收链与五维联合精度尚未建立','CZ-D2数值试验结束，完整性Gate仍FAIL；本轮只编制材料'],subtitle='CZ-D3 · 研究负责人审阅后再决定正式对外交付')
bullets('本次可以交付什么，不能承诺什么',['可以交付：候选机制、条件四参数误差、失败边界及资源缺口','不能把候选跨度、局部信息增益、经验P95混作一种精度','不能把100%数值接受率解释为100%定位正确率','不能写“原单HLA已实现五维联合误差小于10%”'],source='S01/S03/S19')
sl=slide('原单HLA：进展与缺口同时保留','机制证据不替代连续误差与真实接收验证','S03/S04/S05')
ptable(sl,['方向','已有进展','尚缺'],[['机动＋多频','特定15°例候选跨度4 km','同合同连续误差、尾部与阵形'],['公共源幅处理','H1条件不变性','离网格搜索可接受解'],['径向特征','H2速度中位收益','源频/零点、符号与实际提取'],['接收链','CAP06 CAL20覆盖94.3924%','低于95%门；P3未执行']],colwidths=[2,4.6,5.5])
textbox(sl,.7,5.2,12,1,'114,576→2仅为机制验证；最终两标签只差航向4°/5°。\n机动多数已包含在后续实验中，不能重复归因。',18)
sl=slide('原单HLA条件精度：中位收益不等于尾部可靠','P50=场景中位数等权平均；P95=最坏几何无条件尾部','S02/S03')
ptable(sl,['方案','距离%','速度%','输出/计划'],[[r[0],r[1],r[3],r[5]] for r in hrows[:3]],font=18)
textbox(sl,.7,4.5,12,1.9,'每格P50 / P95。H2 U_0.05注入假设径向特征；不是声压提取结果。\nH2每几何32次，尾部探索性强；INF保留拒判。\nAUX-D1 U0是该合同内单节点方位基线，不代表原HLA全部能力。',19)
sl=slide('垂向、倾斜与联合增强：深度空白不能拼接','全部经验测深及五维联合性能仍未建立','S03/S05')
ptable(sl,['资源','可保留证据','限制'],[['短垂向＋4通道','局部深度信息2.34–4.47倍','不是米级P95；校准/环境条件'],['整体倾斜8元','前向误差与数值归因','定位未评价；非物理否定'],['CAP06组合','独立O2/O6控制与接收诊断','融合P3未执行'],['NUM03环境','名义首深折返几何','不是实际可检测性/相干提供器']],colwidths=[2.7,4.5,4.9],font=17)
bullets('CZ-D2：同一双节点合同下检验新面板',['B0=(0,±5) km；B3=(3,±15) km；都不是单HLA','1200 s / 121历元；2 m/s主平台，600 s后15°转向','12个预注册离网格状态；每状态500次、每侧250次','6000基础创新配对8配置；24000公开记录、48000终点'],source='S18/S19',notes='新面板不是对执行者保密的外部盲测。新增目标未用于调参。')
sl=slide('Q1/Q2与P0/P1：输入质量和先验分开说明','有符号方位、同步、关联与偏置支持均为假设','S18/S19')
ptable(sl,['合同','方位σ','导航σ/轴/历元','生成共模 / 半差偏置'],[['Q1','0.05°','25 m','±0.05° / ±0.025°'],['Q2','0.075°','37.5 m','±0.075° / ±0.0375°']],font=18)
textbox(sl,.7,3.7,12,2.6,'P0：估计器采用原支持。P1：仅把估计支持扩大2倍。\n生成偏置不变，P1仍有硬先验，不是超界偏置失配试验。\n归一化尺度改变也可能影响TRF轨迹，称为冻结处理器的支持敏感性。\nQ2同时改变多项条件，不能解释为单一测向噪声因果效应。',19)
sl=slide('新几何Q1/P0：三个参数改善，方位尾部略差','完整性例外下的描述性结果；12几何最坏无条件P95','S08/S09')
sl.shapes.add_picture(str(FIG/'F01_D2_FOUR_PARAMETER_P95.png'),PI(1.85),PI(1.65),height=PI(5.15))
sl=slide('完整八组四参数P95','%用于距离/速度，°用于方位/航向；深度NOT_EVALUATED','S08/S09')
ptable(sl,heads,drows,font=16,y=1.65)
textbox(sl,.65,6.3,12,.5,'每配置6000次；这是经验尾部，不是总体置信保证或工程验收。',16)
sl=slide('速度：池化中位和最坏尾部必须分别报告','旧Q1/B3速度8.566% → 新面板11.557%；两面板并非配对实验','S02/S08')
sl.shapes.add_picture(str(FIG/'F03_SPEED_MEDIAN_AND_TAIL.png'),PI(1.75),PI(1.7),height=PI(5.05))
sl=slide('偏置支持是实质条件，不能省略','B3四组均6000/6000至少一项偏置触界','S08/S10/S18')
sl.shapes.add_picture(str(FIG/'F02_BIAS_SUPPORT_SENSITIVITY.png'),PI(1.7),PI(1.7),height=PI(5.05))
sl=slide('配对变化：保留每个参数的变差记录','Q1/P0；每参数6000对，同一创新并非四份独立证据','S11')
sl.shapes.add_picture(str(FIG/'F04_PAIRED_CHANGES.png'),PI(1.1),PI(1.8),width=PI(11.1))
bullets('完整性例外：原FAIL不改判',['539,763 PASS、1 FAIL；主分析/冷审计各一次，零重试','失败：两份标准库.pyc缓存访问被拒绝，随后读取源代码','日志未观察到成功读取私有真值；不等于OS级隔离认证','负责人接受原因解释，未授予整体科研PASS；不再重跑'],source='S13/S14/S15',subtitle='条件性数值结果保留；完整性例外必须伴随精度引用')
sl=slide('资源需求：不是一张已验证的精度排行榜','新增资源的理论潜力、已测配置和工程条件分别记账','S03/S04/S12/S18')
ptable(sl,['路径','可能受益状态','新增条件与未闭合问题'],[['原阵接收链','距离/运动候选','复数据、阵形、AGC、源频参考、消歧'],['短垂向','局部深度信息','＋4通道/4 m；差分增益、可信传播'],['辅助节点','距离/速度/航向','5–15.297 km部署；测向、偏置、同步关联'],['联合增强','尚未验证','不能把独立结果拼成五维系统']],colwidths=[2.4,3.2,6.5],font=17)
textbox(sl,.7,5.4,12,.8,'距离落在45–60 km内也不证明第一CZ可检测性；无设备价格、最低数量或经验Pareto承诺。',18)
bullets('交付建议与停止边界',['主线：原单HLA研究进展、具体能力缺口及可提取观测条件','增强参照：双节点四参数结果，连同质量、先验及完整性例外','深度：保留局部信息，不填经验测深P95，不承诺五维精度','本轮提交后STOP；等待对外措辞审阅，不发甲方、不启动新实验'],source='S01/S03/S15/S18',subtitle='研究创新落在观测可提取性、不可辨识方向与资源条件之间的联系')
prs.save(OUT/'CZ_D3_INTERNAL_BRIEFING_DRAFT.pptx')
savecsv('PPT_SLIDE_EVIDENCE.csv',slide_index)
write(OUT/'PPT_SPEAKER_NOTES.md','# PPT逐页证据与讲解限制\n\n'+BANNER+'\n\n'+'\n'.join(f"## {r['slide']}. {r['title']}\n\n来源：{r['source_ids']}。{r['qualification']}\n\n不得省略页脚的内部稿及完整性FAIL声明。\n" for r in slide_index))
write(OUT/'RELEASE_REVIEW_CHECKLIST.md','''# 对外封版前审阅清单

状态：INTERNAL_DRAFT_NOT_APPROVED_FOR_EXTERNAL_RELEASE

- [ ] 负责人接受单HLA主线、增强方案附表及深度NOT_EVALUATED表述。
- [ ] 每项精度连同输入Q/P合同、分母、聚合口径和经验统计边界。
- [ ] CZ-D2所有精度引用同时披露539,763 PASS、1 FAIL及整体科研PASS未授予。
- [ ] 未将私有成功读取未观察到升级为完整隔离认证。
- [ ] 未把接受率当正确率、候选跨度当估计P95、局部信息当米级测深。
- [ ] 未把双节点或假设特征能力写成原单HLA实测能力。
- [ ] 未承诺五维联合10%、真实第一CZ覆盖或未经验证的硬件最低配置。
- [ ] 负责人批准最终接收方、版本与发送范围后，才另行授权对外发送。

本清单不授权任何新实验。main、历史Gate、R4=0%保持。
''')
write(OUT/'GPT_SYNC_CZ_D3.md',f'''# CZ-D3内部编稿交付

{BANNER}

父提交：`{PARENT}`；工作分支：`research/cz-d3`。最终SHA由本交付提交封存并回传。

已编制内部技术报告（Markdown、Word）、15页PPT草稿、4组科研图（PNG/PDF/SVG）、统一精度条件表、逐场景/分侧及配对表、架构资源矩阵、来源SHA索引与逐条措辞证据。统计展示均来自存档，不运行新传播、优化、Monte Carlo、Fisher或冷审计。

主线保持原单HLA；双节点为条件性增强参照。H2中位收益与INF尾部、CAP06接收门失败、H3局部信息及T2未评估均保留。CZ-D2仍为539,763 PASS、1 FAIL、integrity_pass=false；日志未见私有成功读取不构成OS隔离认证，不授予整体科研PASS。

Q1/P0封存的最坏场景速度P95为B0 30.9322%、B3 11.5570%，方位尾部略差；P1和Q2退化及所有B3偏置触界均单列。新旧面板不能作为配对比较，不承诺复现8.566%。

请先审阅 [内部报告](CZ_D3_INTERNAL_TECHNICAL_REPORT.md)、[PPT草稿](CZ_D3_INTERNAL_BRIEFING_DRAFT.pptx) 和 [封版清单](RELEASE_REVIEW_CHECKLIST.md)。材料未批准对外发送，深度与五维精度仍未建立。

提交推送后STOP；等待负责人审阅，不自动开展实验或对外发送。main、历史Gate和R4=0%保持。
''')
print(json.dumps({'report_sections':len(sections),'slides':len(prs.slides),'precision_rows':len(precision),'source_files':len(sources),'table_checks':len(checks),'new_scientific_runs':0}))