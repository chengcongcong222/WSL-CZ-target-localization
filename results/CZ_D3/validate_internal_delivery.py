"""Publishing validation only; never imports an experiment or its cold audit."""
from pathlib import Path
import csv,json,hashlib,zipfile,subprocess,re
from pptx import Presentation
from docx import Document
ROOT=Path(__file__).resolve().parents[2];OUT=Path(__file__).resolve().parent
PARENT='8147352b50168ce0f747aa2e42801bf7a2d13b35'
def rows(p):
 with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
checks=[]
def check(name,ok,detail=''):checks.append({'check':name,'pass':bool(ok),'detail':detail})
for r in rows(OUT/'SOURCE_INDEX.csv'):
 p=ROOT/r['path'];current=p.read_bytes();historical=subprocess.check_output(['git','show',PARENT+':'+r['path']])
 check('source_identity_'+r['id'],hashlib.sha256(current).hexdigest()==r['sha256'] and current==historical,r['path'])
a=json.loads((ROOT/'results/CZ_D2/COMMIT_B_RESUME_01/COLD_AUDIT.json').read_text())
check('original_cold_failure_retained',a['PASS_count']==539763 and a['FAIL_count']==1 and a['PASS'] is False)
check('parent_unchanged',subprocess.check_output(['git','rev-parse','HEAD'],text=True).strip()==PARENT)
check('tracked_history_unchanged',subprocess.check_output(['git','diff','HEAD','--name-only'],text=True).strip()=='')
check('saved_table_max_identity_32',all(c['pass'] for c in json.loads((OUT/'TABLE_CHECKS.json').read_text())['checks']))
p=Presentation(OUT/'CZ_D3_INTERNAL_BRIEFING_DRAFT.pptx')
check('ppt_15_slides',len(p.slides)==15)
texts=[]
for i,s in enumerate(p.slides,1):
 text='\n'.join(x.text for x in s.shapes if x.has_text_frame)
 texts.append(text)
 check('slide_footer_'+str(i),'未批准对外发送' in text and 'FAIL' in text)
 check('slide_bounds_'+str(i),all(x.left>=0 and x.top>=0 and x.left+x.width<=p.slide_width and x.top+x.height<=p.slide_height for x in s.shapes))
 check('slide_sources_'+str(i),'SOURCE_INDEX.csv' in s.notes_slide.notes_text_frame.text)
# Compare displayed four-parameter P95 table against archived aggregate values.
agg={(r['config'],r['metric']):r for r in rows(ROOT/'results/CZ_D2/COMMIT_B_RESUME_01/AGGREGATES.csv')}
slide_table=next(x.table for x in p.slides[8].shapes if x.has_table)
metrics=[('range_rel',100),('bearing_circular_deg',1),('speed_rel',100),('heading_circular_deg',1)]
for rr in list(slide_table.rows)[1:]:
 c=rr.cells[0].text
 for j,(m,f) in enumerate(metrics,1):check('ppt_numeric_'+c+'_'+m,rr.cells[j].text==f"{float(agg[c,m]['worst_scene_P95'])*f:.3f}")
report=(OUT/'CZ_D3_INTERNAL_TECHNICAL_REPORT.md').read_text(encoding='utf-8')
for target in re.findall(r'\]\(([^)]+)\)',report):
 if not target.startswith('http'):check('report_link', (OUT/target).exists(),target)
for phrase in ['539,763 PASS、1 FAIL','integrity_pass=false','NOT_EVALUATED','不是操作系统级沙箱','不予授予','不是实际超界偏置失配试验']:
 check('mandatory_qualification',phrase in report,phrase)
doc=Document(OUT/'CZ_D3_INTERNAL_TECHNICAL_REPORT.docx')
check('docx_readable',len(doc.paragraphs)>50 and len(doc.tables)>=7)
check('docx_internal_header','未批准对外发送' in doc.sections[0].header.paragraphs[0].text)
check('precision_rows',len(rows(OUT/'PRECISION_OBSERVATION_CONDITIONS.csv'))==60)
check('depth_not_evaluated',all(r['depth']=='NOT_EVALUATED' for r in rows(OUT/'PRECISION_OBSERVATION_CONDITIONS.csv')))
check('ppt_rendered',len(list((OUT/'slide_previews').glob('*.PNG')))==15 and (OUT/'CZ_D3_INTERNAL_BRIEFING_DRAFT.pdf').stat().st_size>100000)
for fn in ['CZ_D3_INTERNAL_BRIEFING_DRAFT.pptx','CZ_D3_INTERNAL_TECHNICAL_REPORT.docx']:
 with zipfile.ZipFile(OUT/fn) as z:check('ooxml_zip_'+fn,z.testzip() is None)
# Refresh contact sheet from revised rendered slides.
from PIL import Image,ImageDraw
files=sorted((OUT/'slide_previews').glob('*.PNG'),key=lambda p:int(''.join(filter(str.isdigit,p.stem))))
sheet=Image.new('RGB',(1600,1600),'#cccccc')
for i,f in enumerate(files):
 im=Image.open(f);im.thumbnail((520,292));x=(i%3)*533;y=(i//3)*320
 sheet.paste(im,(x,y));ImageDraw.Draw(sheet).text((x+4,y+293),str(i+1),fill='black')
sheet.save(OUT/'slide_previews/CONTACT_SHEET.jpg')
result={'scope':'EDITORIAL_TABLE_AND_FILE_VALIDATION_NOT_SCIENTIFIC_AUDIT','pass':all(c['pass'] for c in checks),'checks':checks,'new_scientific_runs':0,'original_integrity_pass':False,'external_release_authorized':False,'visual_review':'15-slide PowerPoint rendered contact sheet reviewed; corrected two image overflows; final slide bounds checked','word_pdf':'NOT_DELIVERED_BACKGROUND_EXPORT_DID_NOT_COMPLETE; DOCX structure checked, not page-render certified'}
(OUT/'EDITORIAL_VALIDATION.json').write_text(json.dumps(result,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
print(json.dumps({'pass':result['pass'],'checks':len(checks),'failures':[c for c in checks if not c['pass']]},ensure_ascii=False))
assert result['pass']